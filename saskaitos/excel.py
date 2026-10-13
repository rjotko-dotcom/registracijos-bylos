"""Duomenų įrašymas į Excel (.xlsx) – tvarkingas sąskaitų registras."""

from __future__ import annotations

import shutil
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .parser import Invoice

SHEET = "Sąskaitos"
OLD_SHEET = "Saskaitos"  # ankstesnės versijos formatas – konvertuojamas automatiškai
TABLE = "SaskaituRegistras"
HEADER_ROW = 4
FIRST_ROW = HEADER_ROW + 1
ASSETS = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent.parent)) / "saskaitos" / "assets"

BLACK, INK, MUTED, RED = "000000", "1A1A1A", "6B6B6B", "C3002F"
FONT = "Calibri"
MONO = "Consolas"

# (antraštė, Invoice laukas, plotis, lygiavimas, skaičių formatas, šriftas)
COLUMNS = [
    ("Invoice No.",     "invoice_number", 14, "center", "@",          FONT),
    ("Invoice Date",    "invoice_date",   14, "center", "yyyy-mm-dd", FONT),
    ("VIN",             "vin",            22, "center", "@",          MONO),
    ("Kerb Weight, kg", "kerb_weight",    16, "right",  "#,##0",      FONT),
    ("CO₂, g/km",       "co2",            12, "right",  "0.0",        FONT),
    ("Color",           "color",          22, "left",   "@",          FONT),
    ("Decor",           "decor",          44, "left",   "@",          FONT),
    ("Factory Code",    "factory_code",   24, "left",   "@",          MONO),
    ("Total, €",        "total",          15, "right",  "#,##0.00",   FONT),
]
HEADERS = [c[0] for c in COLUMNS]
LAST_COL = get_column_letter(len(COLUMNS))
TOTAL_COL = get_column_letter(len(COLUMNS))

# Įspėjimai iš nuskaitymo priskiriami konkrečiam langeliui
WARNING_FIELDS = {"VIN": "vin", "data": "invoice_date", "Kerb": "kerb_weight", "CO2": "co2"}

MISSING_FILL = PatternFill("solid", fgColor="FDECEF")
WARNING_FILL = PatternFill("solid", fgColor="FFF4E0")
ROW_BORDER = Border(bottom=Side(style="thin", color="E3E3E3"))


class ExcelLockedError(Exception):
    """Excel failas atidarytas kitoje programoje ir negali būti įrašytas."""


# ---------------------------------------------------------------- struktūra

def _new_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110

    for i, (_, _, width, *_rest) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Antraštės blokas: logotipas, pavadinimas, suvestinė
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 6
    logo = ASSETS / "nissan_logo_black.png"
    if logo.exists():
        img = XLImage(str(logo))
        img.height, img.width = 52, round(52 * img.width / img.height)
        ws.add_image(img, "A1")
    ws["B1"] = "SĄSKAITŲ REGISTRAS"
    ws["B1"].font = Font(name=FONT, size=18, bold=True, color=BLACK)
    ws["B1"].alignment = Alignment(vertical="bottom")
    ws["B2"] = "Duomenys automatiškai perkelti iš sąskaitų faktūrų (PDF)"
    ws["B2"].font = Font(name=FONT, size=9, color=MUTED)
    ws["B2"].alignment = Alignment(vertical="top")

    label_col, value_col = get_column_letter(len(COLUMNS) - 1), TOTAL_COL
    summary = [
        (1, "Sąskaitų", f"=SUBTOTAL(103,A{FIRST_ROW}:A1048576)", "0"),
        (2, "Bendra suma, €", f"=SUBTOTAL(109,{TOTAL_COL}{FIRST_ROW}:{TOTAL_COL}1048576)", "#,##0.00"),
    ]
    for row, label, formula, fmt in summary:
        lc, vc = ws[f"{label_col}{row}"], ws[f"{value_col}{row}"]
        lc.value, vc.value = label, formula
        lc.font = Font(name=FONT, size=9, color=MUTED)
        lc.alignment = Alignment(horizontal="right", vertical="center")
        vc.font = Font(name=FONT, size=12, bold=True, color=BLACK)
        vc.alignment = Alignment(horizontal="right", vertical="center")
        vc.number_format = fmt

    # Nissan raudona akcento linija
    for col in range(1, len(COLUMNS) + 1):
        ws.cell(3, col).fill = PatternFill("solid", fgColor=RED)

    # Lentelės antraštė
    ws.row_dimensions[HEADER_ROW].height = 30
    for col, (title, _key, _w, align, *_r) in enumerate(COLUMNS, start=1):
        cell = ws.cell(HEADER_ROW, col, title)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLACK)
        cell.alignment = Alignment(horizontal=align, vertical="center", indent=1 if align != "center" else 0)

    ws.freeze_panes = f"A{FIRST_ROW}"

    # Spausdinimas: gulsčias lapas, visi stulpeliai viename puslapyje
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.left.text = "Sąskaitų registras"
    ws.oddFooter.right.text = "&P / &N"
    wb.properties.title = "Sąskaitų registras"
    return wb


def _find_header_row(ws) -> int | None:
    for row in range(1, 11):
        if ws.cell(row, 1).value in ("Invoice No.", "Invoice number"):
            return row
    return None


def _data_rows(ws) -> list[tuple]:
    header = _find_header_row(ws)
    if header is None:
        return []
    rows = []
    for values in ws.iter_rows(min_row=header + 1, max_col=len(COLUMNS), values_only=True):
        if any(v not in (None, "") for v in values):
            rows.append(values)
    return rows


def _migrate(path: Path, wb) -> Workbook:
    """Ankstesnės versijos failą (su Failas/Įkelta/Pastabos) perrašo nauju formatu."""
    old_rows = _data_rows(wb[OLD_SHEET])
    backup = path.with_name(f"{path.stem} (senas formatas){path.suffix}")
    if not backup.exists():
        shutil.copy2(path, backup)
    new = _new_workbook()
    ws = new[SHEET]
    for values in old_rows:
        _write_row(ws, dict(zip((c[1] for c in COLUMNS), values)), notes={})
    _resize_table(ws)
    return new


def _open(path: Path) -> Workbook:
    if not path.exists():
        return _new_workbook()
    wb = load_workbook(path)
    if SHEET in wb.sheetnames:
        return wb
    if OLD_SHEET in wb.sheetnames:
        return _migrate(path, wb)
    raise ValueError(f"Faile {path.name} nėra lapo „{SHEET}“.")


# ---------------------------------------------------------------- įrašymas

def _last_row(ws) -> int:
    row = ws.max_row
    while row >= FIRST_ROW and all(ws.cell(row, c).value in (None, "") for c in range(1, len(COLUMNS) + 1)):
        row -= 1
    return max(row, HEADER_ROW)


def _write_row(ws, values: dict, notes: dict[str, tuple[str, PatternFill]]) -> None:
    row = _last_row(ws) + 1
    ws.row_dimensions[row].height = 22
    for col, (_title, key, _w, align, fmt, font) in enumerate(COLUMNS, start=1):
        value = values.get(key)
        if isinstance(value, datetime):
            value = value.date()
        cell = ws.cell(row, col, value)
        cell.number_format = fmt if not isinstance(value, date) else "yyyy-mm-dd"
        cell.font = Font(name=font, size=10, color=INK)
        cell.alignment = Alignment(horizontal=align, vertical="center", indent=1 if align != "center" else 0)
        cell.border = ROW_BORDER
        if key in notes:
            text, fill = notes[key]
            cell.fill = fill
            cell.comment = Comment(text, "Sąskaitos → Excel", width=260, height=60)


def _resize_table(ws) -> None:
    ref = f"A{HEADER_ROW}:{LAST_COL}{max(_last_row(ws), FIRST_ROW)}"
    if TABLE in ws.tables:
        ws.tables[TABLE].ref = ref
        ws.tables[TABLE].autoFilter = AutoFilter(ref=ref)
    else:
        table = Table(displayName=TABLE, ref=ref, autoFilter=AutoFilter(ref=ref))
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(table)


def _notes_for(inv: Invoice) -> dict[str, tuple[str, PatternFill]]:
    notes: dict[str, tuple[str, PatternFill]] = {}
    for _title, key, *_ in COLUMNS:
        if getattr(inv, key) in (None, ""):
            notes[key] = (f"Sąskaitoje nerasta ({inv.source_file})", MISSING_FILL)
    for warning in inv.warnings:
        key = next((k for word, k in WARNING_FIELDS.items() if word in warning), None)
        if key:
            notes[key] = (warning, WARNING_FILL)
    return notes


def existing_invoice_numbers(path: Path) -> set[str]:
    if not path.exists():
        return set()
    wb = load_workbook(path, read_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb[OLD_SHEET] if OLD_SHEET in wb.sheetnames else None
    if ws is None:
        return set()
    numbers = set()
    header_seen = False
    for (value,) in ws.iter_rows(max_col=1, values_only=True):
        if value in ("Invoice No.", "Invoice number"):
            header_seen = True
        elif header_seen and value not in (None, ""):
            numbers.add(str(value).strip())
    wb.close()
    return numbers


def append_invoices(path: Path, invoices: list[Invoice]) -> None:
    """Prideda eilutes ir išsaugo. Rašoma į laikiną failą, kad nesugadintų esamo."""
    if not invoices:
        return
    wb = _open(path)
    ws = wb[SHEET]
    for inv in invoices:
        _write_row(ws, {key: getattr(inv, key) for _t, key, *_ in COLUMNS}, _notes_for(inv))
    _resize_table(ws)

    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f"~tmp_{path.name}")
    try:
        wb.save(tmp)
        tmp.replace(path)
    except PermissionError as exc:
        tmp.unlink(missing_ok=True)
        raise ExcelLockedError(
            f"Nepavyko įrašyti į {path.name} – uždarykite jį Excel'yje ir bandykite dar kartą."
        ) from exc
